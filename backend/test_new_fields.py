import io
from datetime import date
import openpyxl
from app.excel_generator import build_daily_excel_report
from app.main import compute_data_quality_flag

def test_quality_flag():
    flag_ok, note_ok = compute_data_quality_flag(
        total_hh=10000,
        segregated_hh=9200,
        generation_mt=15.0,
        wet_proc_mt=8.0,
        dry_proc_mt=6.0,
        total_cap_mt=20.0,
        dump_yard_mt=1.0
    )
    print(f"Normal Case Flag: {flag_ok} | Note: {note_ok}")
    assert flag_ok == "OK"

    flag_err, note_err = compute_data_quality_flag(
        total_hh=10000,
        segregated_hh=12500,  # 12500 > 10000 -> Exceeded!
        generation_mt=15.0,
        wet_proc_mt=8.0,
        dry_proc_mt=6.0,
        total_cap_mt=20.0,
        dump_yard_mt=1.0
    )
    print(f"Exceeded Case Flag: {flag_err} | Note: {note_err}")
    assert flag_err == "Segregated HH>Total HH"

def test_excel_export():
    sample_data = {
        "Corporation": [
            {
                "s_no": 1,
                "name": "Coimbatore",
                "total_households": 560000,
                "segregated_hh_collected": 510000,
                "total_generation_today_mt": 1050.0,
                "wet_waste_processed_mt": 620.0,
                "compost_output_mt": 95.0,
                "dry_facilities_capacity_mt": 400.0,
                "recyclable_sold_mt": 45.0,
                "dry_waste_cement_mt": 180.0,
                "dump_yard_mt": 205.0,
                "data_quality_flag": "VALID"
            }
        ],
        "Chengalpattu Region": [
            {
                "s_no": 1,
                "name": "Chengalpattu",
                "total_households": 32000,
                "segregated_hh_collected": 28500,
                "total_generation_today_mt": 42.0,
                "wet_waste_processed_mt": 24.0,
                "compost_output_mt": 3.8,
                "dry_facilities_capacity_mt": 18.0,
                "recyclable_sold_mt": 2.2,
                "dry_waste_cement_mt": 8.5,
                "dump_yard_mt": 7.3,
                "data_quality_flag": "VALID"
            }
        ]
    }
    buf = build_daily_excel_report(date.today(), sample_data)
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    print(f"Daily Excel generated successfully. Sheet title: {ws.title}, max_col: {ws.max_column}, max_row: {ws.max_row}")
    assert ws.max_column == 15
    assert ws["A3"].value == "S.No"
    assert ws["B3"].value == "Name of the Corporation / Municipality"
    assert ws["C4"].value == "No. of HHs"
    assert ws["D5"].value == "No. of HHs segregated waste collected"
    assert ws["E5"].value == "% of Collection"
    assert ws["F3"].value == "Total waste generation Today (MT)"
    assert ws["G4"].value == "Wet waste processed MCC / BIO-METHA (MT)"
    assert ws["H4"].value == "Output quantity as compost per day (MT)"
    assert ws["I4"].value == "Processing facilities available (MT)"
    assert ws["J4"].value == "Recyclable waste sold by sanitary workers (MT)"
    assert ws["K4"].value == "Dry waste disposed (Cement industries / recycling units) (MT)"
    assert ws["L4"].value == "Total processed waste (MT)"
    assert ws["M4"].value == "Waste sent to dumping yard (MT)"
    assert ws["N4"].value == "% of processed waste"
    assert ws["O4"].value == "Data-quality flag"
    print("All 15 Daily Columns verified successfully!")

def test_monthly_excel_export():
    from app.excel_generator import build_monthly_excel_report
    sample_monthly = {
        "Corporation": [
            {
                "s_no": 1,
                "name": "Coimbatore",
                "total_households": 560000,
                "segregated_hh_collected": 510000,
                "total_generation_today_mt": 31500.0,
                "wet_waste_processed_mt": 18600.0,
                "compost_output_mt": 2850.0,
                "dry_facilities_capacity_mt": 400.0,
                "recyclable_sold_mt": 1350.0,
                "dry_waste_cement_mt": 5400.0,
                "dump_yard_mt": 6150.0,
                "data_quality_flag": "OK",
                "days_reported": 30
            }
        ],
        "Chengalpattu Region": [
            {
                "s_no": 1,
                "name": "Chengalpattu",
                "total_households": 32000,
                "segregated_hh_collected": 28500,
                "total_generation_today_mt": 1260.0,
                "wet_waste_processed_mt": 720.0,
                "compost_output_mt": 114.0,
                "dry_facilities_capacity_mt": 18.0,
                "recyclable_sold_mt": 66.0,
                "dry_waste_cement_mt": 255.0,
                "dump_yard_mt": 219.0,
                "data_quality_flag": "OK",
                "days_reported": 30
            }
        ]
    }
    buf = build_monthly_excel_report(2026, 9, sample_monthly)
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    print(f"Monthly Excel generated successfully. Sheet title: {ws.title}, max_col: {ws.max_column}, max_row: {ws.max_row}")
    assert ws.max_column == 15
    assert ws["A3"].value == "S.No"
    assert ws["B3"].value == "Name of the Corporation / Municipality"
    assert ws["C4"].value == "No. of HHs"
    assert ws["D4"].value == "Door to Door Collection (Monthly Avg)"
    assert ws["D5"].value == "No. of HHs segregated waste collected"
    assert ws["E5"].value == "% of Collection"
    assert ws["F3"].value == "Total waste generation (MT)"
    assert ws["G4"].value == "Wet waste processed MCC / BIO-METHA (MT)"
    assert ws["H4"].value == "Output quantity as compost (MT)"
    assert ws["I4"].value == "Processing facilities available (MT)"
    assert ws["J4"].value == "Recyclable waste sold by sanitary workers (MT)"
    assert ws["K4"].value == "Dry waste disposed (Cement industries / recycling units) (MT)"
    assert ws["L4"].value == "Total processed waste (MT)"
    assert ws["M4"].value == "Waste sent to dumping yard (MT)"
    assert ws["N4"].value == "% of processed waste"
    assert ws["O4"].value == "Data-quality flag"
    print("All 15 Monthly Columns verified successfully!")

if __name__ == "__main__":
    test_quality_flag()
    test_excel_export()
    test_monthly_excel_export()
    print("ALL BACKEND & EXCEL TESTS PASSED SUCCESSFULLY!")
