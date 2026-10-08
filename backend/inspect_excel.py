import openpyxl

file_path = r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n==========================================")
    print(f"Sheet: {sheetname} (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"==========================================")
    for r in range(1, min(35, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else "" for c in range(1, min(20, ws.max_column + 1))]
        if any(row_vals):
            print(f"R{r:02d}: " + " | ".join(row_vals[:16]))
