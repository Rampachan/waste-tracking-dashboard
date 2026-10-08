import openpyxl
import json

file_path = r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['ULB wise Format-22.05.2026']

ulbs = []
current_region = "Corporation"

def to_num(val):
    if val is None or str(val).strip() in ["", "-"]:
        return 0.0
    try:
        return float(val)
    except:
        return 0.0

for r in range(5, ws.max_row + 1):
    sno = ws.cell(r, 1).value
    name = ws.cell(r, 2).value

    sno_str = str(sno).strip() if sno is not None else ""
    name_str = str(name).strip() if name is not None else ""

    if not name_str and not sno_str:
        continue

    if ("Region" in name_str or "region" in name_str) and not sno_str.isdigit():
        current_region = name_str
        continue

    if sno_str.isdigit() or (name_str and name_str not in ["Total", "TOTAL", "Region Total"]):
        hh = int(to_num(ws.cell(r, 17).value))
        mcc = to_num(ws.cell(r, 20).value)
        mrf = to_num(ws.cell(r, 22).value)
        bio = to_num(ws.cell(r, 24).value)
        oth = to_num(ws.cell(r, 26).value)
        pw = int(to_num(ws.cell(r, 3).value))
        ow = int(to_num(ws.cell(r, 6).value))
        pv = int(to_num(ws.cell(r, 9).value))
        sv = int(to_num(ws.cell(r, 13).value))

        cat = "Corporation" if current_region == "Corporation" else "Municipality"

        ulbs.append({
            "s_no": int(sno_str),
            "region": current_region,
            "category": cat,
            "name": name_str,
            "households": hh,
            "mcc": mcc,
            "mrf": mrf,
            "biometh": bio,
            "other": oth,
            "perm_workers": pw,
            "outsource_workers": ow,
            "primary_veh": pv,
            "secondary_veh": sv
        })

print(f"Loaded {len(ulbs)} ULBs with official data.")
with open(r"C:\Users\Admin\.gemini\antigravity\scratch\waste-tracking-dashboard\backend\app\official_ulb_static_data.json", "w", encoding="utf-8") as f:
    json.dump(ulbs, f, indent=2)
print("Saved official static data to backend/app/official_ulb_static_data.json")
