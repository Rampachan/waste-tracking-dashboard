import openpyxl

file_path = r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['ULB wise Format-22.05.2026']

print("=== Scanning all rows of ULB wise Format-22.05.2026 ===")
categories = []
current_cat = "Corporation"

ulb_static_data = []

for r in range(5, ws.max_row + 1):
    sno = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    
    if name is None and sno is None:
        continue

    sno_str = str(sno).strip() if sno is not None else ""
    name_str = str(name).strip() if name is not None else ""

    # Check if banner / region header
    if ("Region" in name_str or "region" in name_str or "TOTAL" in name_str.upper()) and not sno_str.isdigit():
        print(f"Row {r:3d}: BANNER/TOTAL -> '{sno_str}' | '{name_str}'")
        if "Region" in name_str or "region" in name_str:
            current_cat = name_str
        continue

    # ULB data row
    if sno_str.isdigit() or (name_str and name_str != "Total" and name_str != "TOTAL"):
        # Extract the static columns
        perm_in_pos = ws.cell(r, 3).value    # Col C: Permanent workers in position
        outsource_agr = ws.cell(r, 6).value  # Col F: Outsourcing agreement
        pri_veh_allot = ws.cell(r, 9).value  # Col I: Primary vehicles allotted
        sec_veh_allot = ws.cell(r, 13).value # Col M: Secondary vehicles allotted
        households = ws.cell(r, 17).value    # Col Q: Households
        mcc_cap = ws.cell(r, 20).value       # Col T: MCC Capacity
        mrf_cap = ws.cell(r, 22).value       # Col V: MRF Capacity
        bio_cap = ws.cell(r, 24).value       # Col X: BioMeth Capacity
        oth_cap = ws.cell(r, 26).value       # Col Z: Other facility Capacity

        ulb_static_data.append({
            "row": r,
            "region": current_cat,
            "s_no": sno_str,
            "name": name_str,
            "perm_workers": perm_in_pos,
            "outsource_workers": outsource_agr,
            "primary_veh": pri_veh_allot,
            "secondary_veh": sec_veh_allot,
            "households": households,
            "mcc_cap": mcc_cap,
            "mrf_cap": mrf_cap,
            "bio_cap": bio_cap,
            "oth_cap": oth_cap
        })

print(f"\nTotal ULBs extracted: {len(ulb_static_data)}")
print("\nSample 5 extracted ULBs:")
for u in ulb_static_data[:5]:
    print(u)

print("\nLast 5 extracted ULBs:")
for u in ulb_static_data[-5:]:
    print(u)
