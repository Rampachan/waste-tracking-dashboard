import io
from datetime import date
from typing import List, Dict, Any, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Palette
COLOR_HEADER_BG = "1F4E79"       # Deep Navy
COLOR_HEADER_TXT = "FFFFFF"
COLOR_CAT_BANNER = "D9E1F2"      # Soft Periwinkle for region banners
COLOR_SUBTOTAL_BG = "FFF2CC"     # Soft Gold for subtotals
COLOR_REGION_TOT_BG = "FCE4D6"   # Peach for Region Total
COLOR_GRAND_TOT_BG = "E2EFDA"    # Soft Green for Grand Total

# Status / Flag colors
FLAG_COLORS = {
    "OK": {"fill": "E2EFDA", "font": "276A3C"},
    "Segregated HH>Total HH": {"fill": "FCE4D6", "font": "C00000"},
    "Processed>Generation": {"fill": "FCE4D6", "font": "C00000"},
    "WARNING": {"fill": "FFF2CC", "font": "8C6B00"},
    "PENDING": {"fill": "F2F2F2", "font": "595959"}
}

BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

BORDER_TOTAL = Border(
    top=Side(style='thin', color='000000'),
    bottom=Side(style='double', color='000000'),
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000')
)

def _generate_23_column_swm_report(
    title: str,
    sheet_title: str,
    grouped_data: Dict[str, List[Dict[str, Any]]],
    is_monthly: bool = False
) -> io.BytesIO:
    """
    Builds the official Tamil Nadu SWM Report with the exact 23-column hierarchy:
    Cols 1-5 (A-E): Date, Sl.No., Region, Name of the ULB, No. of HHs
    Col 6 (F): SOURCE SEGREGATION: No. of HHs segregated waste collected
    Cols 7-9 (G-I): WASTE GENERATION: Total waste generation Today (MT), Wet waste generation (MT), Dry waste generation (MT)
    Cols 10-12 (J-L): WET WASTE: Wet waste processed MCC / BIO-METHA (MT) & others, Compost generated today (MT), Manure sold / given at free of cost (MT)
    Cols 13-15 (M-O): DRY WASTE: Recyclable waste sold by sanitary workers (MT), Recyclable waste sent to recyclers (MT), Dry waste disposed to Cement industries (MT)
    Cols 16-17 (P-Q): SANITARY WASTE: Sanitary waste - incineration capacity (MT/day), Sanitary waste incinerated today (MT)
    Col 18 (R): SPECIAL WASTE: Special waste processed / disposed today (MT)
    Cols 19-23 (S-W): CALCULATED - DO NOT ENTER:
      Col 19 (S): Total dry waste processed (MT) =13+14+15
      Col 20 (T): Total processed waste (MT) =10+17+18+19
      Col 21 (U): Waste sent to dumping yard (MT) =7-20
      Col 22 (V): % of processed waste =20/7
      Col 23 (W): Data-quality flag
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    # 1. Main Title (A1:W1)
    ws.merge_cells("A1:W1")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Multi-tier Table Headers (Rows 3, 4, 5)
    # Row 3: Top-level merged groups
    ws.merge_cells("A3:A5")
    ws["A3"] = "Date"
    
    ws.merge_cells("B3:B5")
    ws["B3"] = "Sl.No."

    ws.merge_cells("C3:C5")
    ws["C3"] = "Region"

    ws.merge_cells("D3:D5")
    ws["D3"] = "Name of the ULB"

    ws.merge_cells("E3:E5")
    ws["E3"] = "No. of HHs"

    ws.merge_cells("F3:F4")
    ws["F3"] = "SOURCE SEGREGATION"

    ws.merge_cells("G3:I3")
    ws["G3"] = "WASTE GENERATION"

    ws.merge_cells("J3:L3")
    ws["J3"] = "WET WASTE"

    ws.merge_cells("M3:O3")
    ws["M3"] = "DRY WASTE"

    ws.merge_cells("P3:Q3")
    ws["P3"] = "SANITARY WASTE"

    ws.merge_cells("R3:R4")
    ws["R3"] = "SPECIAL WASTE"

    ws.merge_cells("S3:W3")
    ws["S3"] = "CALCULATED - DO NOT ENTER"

    # Row 4: Intermediate Level
    ws.merge_cells("G4:G5")
    ws["G4"] = "Total waste generation Today (MT)" if not is_monthly else "Total waste generation (MT)"

    ws.merge_cells("H4:H5")
    ws["H4"] = "Wet waste generation (MT)"

    ws.merge_cells("I4:I5")
    ws["I4"] = "Dry waste generation (MT)"

    ws.merge_cells("J4:J5")
    ws["J4"] = "Wet waste processed MCC / BIO-METHA (MT) & others"

    ws.merge_cells("K4:K5")
    ws["K4"] = "Compost generated today (MT)" if not is_monthly else "Compost generated (MT)"

    ws.merge_cells("L4:L5")
    ws["L4"] = "Manure sold / given at free of cost (MT)"

    ws.merge_cells("M4:M5")
    ws["M4"] = "Recyclable waste sold by sanitary workers (MT)"

    ws.merge_cells("N4:N5")
    ws["N4"] = "Recyclable waste sent to recyclers (MT)"

    ws.merge_cells("O4:O5")
    ws["O4"] = "Dry waste disposed to Cement industries (MT)"

    ws.merge_cells("P4:P5")
    ws["P4"] = "Sanitary waste - incineration capacity (MT/day)"

    ws.merge_cells("Q4:Q5")
    ws["Q4"] = "Sanitary waste incinerated today (MT)" if not is_monthly else "Sanitary waste incinerated (MT)"

    ws.merge_cells("S4:S5")
    ws["S4"] = "Total dry waste processed (MT) =13+14+15"

    ws.merge_cells("T4:T5")
    ws["T4"] = "Total processed waste (MT) =10+17+18+19"

    ws.merge_cells("U4:U5")
    ws["U4"] = "Waste sent to dumping yard (MT) =7-20"

    ws.merge_cells("V4:V5")
    ws["V4"] = "% of processed waste =20/7"

    ws.merge_cells("W4:W5")
    ws["W4"] = "Data-quality flag"

    # Row 5: Leaf Headers for F & R
    ws["F5"] = "No. of HHs segregated waste collected"
    ws["R5"] = "Special waste processed / disposed today (MT)"

    # Header Styling
    header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    header_font = Font(name="Calibri", size=9, bold=True, color=COLOR_HEADER_TXT)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in range(3, 6):
        ws.row_dimensions[r].height = 26
        for col_idx in range(1, 24):
            c = ws.cell(row=r, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align
            c.border = Border(
                left=Side(style='thin', color='FFFFFF'),
                right=Side(style='thin', color='FFFFFF'),
                top=Side(style='thin', color='FFFFFF'),
                bottom=Side(style='thin', color='FFFFFF')
            )

    current_row = 6
    region_total_rows = []

    # Section ordering per official master layout
    section_keys = [
        "Corporation",
        "Chengalpattu Region",
        "Salem region",
        "Vellore Region",
        "Tiruppur Region",
        "Madurai Region",
        "Thanjavur region",
        "Tirunelveli Region"
    ]

    for section in section_keys:
        items = grouped_data.get(section, [])
        if not items:
            continue

        # Section Banner Row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=23)
        banner_cell = ws.cell(row=current_row, column=1)
        banner_cell.value = f"{section.upper()}"
        banner_cell.font = Font(name="Calibri", size=10, bold=True, color="1F4E79")
        banner_cell.fill = PatternFill(start_color=COLOR_CAT_BANNER, fill_type="solid")
        banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 22
        current_row += 1

        start_data_row = current_row

        for item in items:
            # Col 1: Date (A)
            ws.cell(row=current_row, column=1, value=str(item.get("log_date", "")))
            # Col 2: Sl.No. (B)
            ws.cell(row=current_row, column=2, value=item.get("s_no", ""))
            # Col 3: Region (C)
            ws.cell(row=current_row, column=3, value=item.get("region", section))
            # Col 4: Name of the ULB (D)
            ws.cell(row=current_row, column=4, value=item.get("name", ""))
            # Col 5: No. of HHs (E)
            ws.cell(row=current_row, column=5, value=item.get("total_households", 0))
            # Col 6: No. of HHs segregated waste collected (F)
            ws.cell(row=current_row, column=6, value=item.get("segregated_hh_collected", 0))

            # Col 7: Total waste generation Today (MT) (G)
            ws.cell(row=current_row, column=7, value=float(item.get("total_generation_today_mt") or 0.0))
            # Col 8: Wet waste generation (MT) (H)
            ws.cell(row=current_row, column=8, value=float(item.get("wet_waste_generation_mt") or 0.0))
            # Col 9: Dry waste generation (MT) (I)
            ws.cell(row=current_row, column=9, value=float(item.get("dry_waste_generation_mt") or 0.0))

            # Col 10: Wet waste processed MCC / BIO-METHA (MT) & others (J)
            ws.cell(row=current_row, column=10, value=float(item.get("wet_waste_processed_mt") or 0.0))
            # Col 11: Compost generated today (MT) (K)
            ws.cell(row=current_row, column=11, value=float(item.get("compost_output_mt") or 0.0))
            # Col 12: Manure sold / given at free of cost (MT) (L)
            ws.cell(row=current_row, column=12, value=float(item.get("manure_sold_free_mt") or 0.0))

            # Col 13: Recyclable waste sold by sanitary workers (MT) (M)
            ws.cell(row=current_row, column=13, value=float(item.get("recyclable_sold_mt") or 0.0))
            # Col 14: Recyclable waste sent to recyclers (MT) (N)
            ws.cell(row=current_row, column=14, value=float(item.get("recyclable_sent_to_recyclers_mt") or 0.0))
            # Col 15: Dry waste disposed to Cement industries (MT) (O)
            ws.cell(row=current_row, column=15, value=float(item.get("dry_waste_cement_mt") or 0.0))

            # Col 16: Sanitary waste - incineration capacity (MT/day) (P)
            ws.cell(row=current_row, column=16, value=float(item.get("sanitary_waste_incineration_capacity_mt") or 0.0))
            # Col 17: Sanitary waste incinerated today (MT) (Q)
            ws.cell(row=current_row, column=17, value=float(item.get("sanitary_waste_incinerated_today_mt") or 0.0))

            # Col 18: Special waste processed / disposed today (MT) (R)
            ws.cell(row=current_row, column=18, value=float(item.get("special_waste_processed_mt") or 0.0))

            # Col 19: Total dry waste processed (MT) =13+14+15 (S)
            ws.cell(row=current_row, column=19, value=f"=M{current_row}+N{current_row}+O{current_row}")
            # Col 20: Total processed waste (MT) =10+17+18+19 (T)
            ws.cell(row=current_row, column=20, value=f"=J{current_row}+Q{current_row}+R{current_row}+S{current_row}")
            # Col 21: Waste sent to dumping yard (MT) =7-20 (U)
            ws.cell(row=current_row, column=21, value=f"=G{current_row}-T{current_row}")
            # Col 22: % of processed waste =20/7 (V)
            ws.cell(row=current_row, column=22, value=f"=IF(G{current_row}>0, T{current_row}/G{current_row}, 0)")
            # Col 23: Data-quality flag (W)
            ws.cell(row=current_row, column=23, value=f'=IF(F{current_row}>E{current_row}, "Segregated HH>Total HH", IF(T{current_row}>G{current_row}, "Processed>Generation", "OK"))')

            # Formatting Data Row
            for c_idx in range(1, 24):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.font = Font(name="Calibri", size=9)
                cell.border = BORDER_THIN

                # Alignments and number formats
                if c_idx in [1, 2, 3]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 4:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx in [5, 6]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                elif c_idx in range(7, 22):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0.0"
                elif c_idx == 22: # % Col V
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.0%"
                elif c_idx == 23: # Flag Col W
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        end_data_row = current_row - 1

        # Region Subtotal Row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        sub_title = ws.cell(row=current_row, column=1)
        sub_title.value = f"{section} Subtotal"
        sub_title.font = Font(name="Calibri", size=9, bold=True, color="1F4E79")
        sub_title.alignment = Alignment(horizontal="right", vertical="center")

        # Sum Col E (HHs) & Col F (Seg HHs)
        ws.cell(row=current_row, column=5, value=f"=SUM(E{start_data_row}:E{end_data_row})").number_format = "#,##0"
        ws.cell(row=current_row, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})").number_format = "#,##0"

        # Sum Cols G to R (MT)
        for c_idx in range(7, 19):
            col_letter = get_column_letter(c_idx)
            cell = ws.cell(row=current_row, column=c_idx, value=f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})")
            cell.number_format = "#,##0.0"

        # Calculated formulas for subtotals
        ws.cell(row=current_row, column=19, value=f"=M{current_row}+N{current_row}+O{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=20, value=f"=J{current_row}+Q{current_row}+R{current_row}+S{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=21, value=f"=G{current_row}-T{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=22, value=f"=IF(G{current_row}>0, T{current_row}/G{current_row}, 0)").number_format = "0.0%"
        ws.cell(row=current_row, column=23, value=f'=IF(F{current_row}>E{current_row}, "Segregated HH>Total HH", IF(T{current_row}>G{current_row}, "Processed>Generation", "OK"))').alignment = Alignment(horizontal="center")

        for c_idx in range(1, 24):
            c = ws.cell(row=current_row, column=c_idx)
            c.font = Font(name="Calibri", size=9, bold=True, color="1F4E79")
            c.fill = PatternFill(start_color=COLOR_SUBTOTAL_BG, fill_type="solid")
            c.border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
            if c_idx >= 5 and c_idx <= 22:
                c.alignment = Alignment(horizontal="right", vertical="center")

        region_total_rows.append(current_row)
        ws.row_dimensions[current_row].height = 22
        current_row += 1

    # Grand Total Row
    if region_total_rows:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        gt_title = ws.cell(row=current_row, column=1)
        gt_title.value = "STATEWIDE GRAND TOTAL (169 ULBs)"
        gt_title.font = Font(name="Calibri", size=10, bold=True, color="1F4E79")
        gt_title.alignment = Alignment(horizontal="right", vertical="center")

        # Sum of region totals
        def build_sum_formula(col_idx):
            col_letter = get_column_letter(col_idx)
            return "=" + "+".join(f"{col_letter}{r}" for r in region_total_rows)

        ws.cell(row=current_row, column=5, value=build_sum_formula(5)).number_format = "#,##0"
        ws.cell(row=current_row, column=6, value=build_sum_formula(6)).number_format = "#,##0"

        for c_idx in range(7, 19):
            ws.cell(row=current_row, column=c_idx, value=build_sum_formula(c_idx)).number_format = "#,##0.0"

        # Calculated formulas for grand total
        ws.cell(row=current_row, column=19, value=f"=M{current_row}+N{current_row}+O{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=20, value=f"=J{current_row}+Q{current_row}+R{current_row}+S{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=21, value=f"=G{current_row}-T{current_row}").number_format = "#,##0.0"
        ws.cell(row=current_row, column=22, value=f"=IF(G{current_row}>0, T{current_row}/G{current_row}, 0)").number_format = "0.0%"
        ws.cell(row=current_row, column=23, value=f'=IF(F{current_row}>E{current_row}, "Segregated HH>Total HH", IF(T{current_row}>G{current_row}, "Processed>Generation", "OK"))').alignment = Alignment(horizontal="center")

        for c_idx in range(1, 24):
            c = ws.cell(row=current_row, column=c_idx)
            c.font = Font(name="Calibri", size=10, bold=True, color="1F4E79")
            c.fill = PatternFill(start_color=COLOR_GRAND_TOT_BG, fill_type="solid")
            c.border = BORDER_TOTAL
            if c_idx >= 5 and c_idx <= 22:
                c.alignment = Alignment(horizontal="right", vertical="center")

        ws.row_dimensions[current_row].height = 24

    # Optimal Column Widths
    col_widths = {
        "A": 12,   # Date
        "B": 7,    # Sl.No.
        "C": 18,   # Region
        "D": 26,   # Name of the ULB
        "E": 14,   # No. of HHs
        "F": 22,   # No. of HHs segregated waste collected
        "G": 18,   # Total waste generation Today (MT)
        "H": 16,   # Wet waste generation (MT)
        "I": 16,   # Dry waste generation (MT)
        "J": 22,   # Wet waste processed MCC / BIO-METHA (MT) & others
        "K": 16,   # Compost generated today (MT)
        "L": 18,   # Manure sold / given at free of cost (MT)
        "M": 18,   # Recyclable waste sold by sanitary workers (MT)
        "N": 18,   # Recyclable waste sent to recyclers (MT)
        "O": 18,   # Dry waste disposed to Cement industries (MT)
        "P": 18,   # Sanitary waste - incineration capacity (MT/day)
        "Q": 18,   # Sanitary waste incinerated today (MT)
        "R": 18,   # Special waste processed / disposed today (MT)
        "S": 18,   # Total dry waste processed (MT) =13+14+15
        "T": 18,   # Total processed waste (MT) =10+17+18+19
        "U": 16,   # Waste sent to dumping yard (MT) =7-20
        "V": 14,   # % of processed waste =20/7
        "W": 22    # Data-quality flag
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def build_daily_excel_report(report_date: date, grouped_data: Dict[str, List[Dict[str, Any]]]) -> io.BytesIO:
    """
    Builds the official Tamil Nadu Daily SWM Report with exact 23-column multi-tier hierarchy.
    """
    title = f"TAMIL NADU MUNICIPAL ADMINISTRATION - SOLID WASTE MANAGEMENT DAILY REPORT ({report_date.strftime('%d-%b-%Y')})"
    sheet_title = f"Daily-{report_date.strftime('%d%b%Y')}"
    return _generate_23_column_swm_report(
        title=title,
        sheet_title=sheet_title,
        grouped_data=grouped_data,
        is_monthly=False
    )

def build_monthly_excel_report(year: int, month: int, monthly_data: Union[Dict[str, List[Dict[str, Any]]], List[Dict[str, Any]]]) -> io.BytesIO:
    """
    Builds the official Tamil Nadu Monthly Consolidated SWM Report with exact 23-column multi-tier hierarchy.
    """
    if isinstance(monthly_data, list):
        grouped_data: Dict[str, List[Dict[str, Any]]] = {}
        for item in monthly_data:
            reg = item.get("region", "Corporation")
            if reg not in grouped_data:
                grouped_data[reg] = []
            grouped_data[reg].append(item)
    else:
        grouped_data = monthly_data

    month_date = date(year, month, 1)
    month_name = month_date.strftime('%B')
    title = f"TAMIL NADU MUNICIPAL ADMINISTRATION - SOLID WASTE MANAGEMENT MONTHLY REPORT ({month_name} {year})"
    sheet_title = f"Monthly-{month:02d}_{year}"

    return _generate_23_column_swm_report(
        title=title,
        sheet_title=sheet_title,
        grouped_data=grouped_data,
        is_monthly=True
    )
