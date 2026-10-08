import os
import openpyxl

downloads = r"C:\Users\Admin\Downloads"
for fname in os.listdir(downloads):
    if fname.endswith(".xlsx") and not fname.startswith("~$"):
        fpath = os.path.join(downloads, fname)
        try:
            wb = openpyxl.load_workbook(fpath, read_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                found = False
                for r in ws.iter_rows(max_row=10, values_only=True):
                    for cell in r:
                        if cell and any(keyword in str(cell) for keyword in ["compost", "Recyclable", "Cement", "Data-quality", "segregated"]):
                            print(f"File: {fname} | Sheet: {sname} -> Cell: {cell}")
                            found = True
                            break
                    if found:
                        break
            wb.close()
        except Exception as e:
            pass
