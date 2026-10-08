import openpyxl

file_path = r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

ws = wb['ULB wise Format-22.05.2026']
print(f"Sheet: ULB wise Format-22.05.2026 (rows: {ws.max_row}, cols: {ws.max_column})")

print("\n--- Rows 1 to 4 (Header Hierarchy) ---")
for r in range(1, 5):
    row_vals = [f"Col{c}({openpyxl.utils.get_column_letter(c)}): {ws.cell(r, c).value}" for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
    print(f"Row {r}:", " | ".join(row_vals[:15]))

print("\n--- Full Combined Column Headers (Col 1 to max_col) ---")
for c in range(1, ws.max_column + 1):
    vals = [str(ws.cell(r, c).value).strip() for r in range(1, 5) if ws.cell(r, c).value is not None]
    col_letter = openpyxl.utils.get_column_letter(c)
    print(f"Col {c:2d} ({col_letter}): " + " -> ".join(vals))

print("\n--- Sample First 3 Data Rows (Row 5 to 7) ---")
for r in range(5, 8):
    row_vals = [f"{openpyxl.utils.get_column_letter(c)}: {ws.cell(r, c).value}" for c in range(1, ws.max_column + 1)]
    print(f"Row {r}:", " | ".join(row_vals[:15]))
