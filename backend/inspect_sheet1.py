import openpyxl

file_path = r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
print("Sheet names:", wb.sheetnames)

sheet1_name = wb.sheetnames[0]
ws = wb[sheet1_name]
print(f"\nSheet 1: {sheet1_name} (rows: {ws.max_row}, cols: {ws.max_column})")

print("\n--- Rows 1 to 5 (Headers) ---")
for r in range(1, 6):
    row_vals = [f"Col{c}({openpyxl.utils.get_column_letter(c)}): {ws.cell(r, c).value}" for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
    print(f"Row {r}:", " | ".join(row_vals))

print("\n--- All Column Headers across Rows 1-4 combined ---")
for c in range(1, ws.max_column + 1):
    vals = [str(ws.cell(r, c).value).strip() for r in range(1, 5) if ws.cell(r, c).value is not None]
    col_letter = openpyxl.utils.get_column_letter(c)
    print(f"Col {c:2d} ({col_letter}): " + " -> ".join(vals))
