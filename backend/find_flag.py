import openpyxl, os

d = r"C:\Users\Admin\Downloads"
for f in os.listdir(d):
    if f.endswith(".xlsx") and not f.startswith("~$"):
        fp = os.path.join(d, f)
        try:
            wb = openpyxl.load_workbook(fp, data_only=True)
            for s in wb.sheetnames:
                ws = wb[s]
                for r in range(1, min(10, ws.max_row+1)):
                    for c in range(1, min(40, ws.max_column+1)):
                        val = str(ws.cell(r, c).value or "")
                        if "quality" in val.lower() or "flag" in val.lower() or "segregated hh" in val.lower():
                            print(f"{f} -> {s} -> R{r}C{c}: {val}")
            wb.close()
        except Exception:
            pass
