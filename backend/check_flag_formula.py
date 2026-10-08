import openpyxl

files = [
    r"C:\Users\Admin\Downloads\12.05.2026. - Daily SWM,UWM, Sanitation Format .., (1) (2).xlsx",
    r"C:\Users\Admin\Downloads\SWM - Daily Format  21.05.2026 (1).xlsx"
]

for fpath in files:
    try:
        wb = openpyxl.load_workbook(fpath, data_only=False)
        print("File:", fpath)
        for s in wb.sheetnames:
            ws = wb[s]
            for r in range(1, min(6, ws.max_row+1)):
                for c in range(1, min(ws.max_column+1, 50)):
                    val = ws.cell(r, c).value
                    if val and any(k in str(val).lower() for k in ["flag", "quality", "segregated"]):
                        print(f"Sheet {s} Row {r} Col {c}: {val}")
                        # print sample formula or cell below
                        print(f"  Row 6 Col {c} value/formula: {ws.cell(6, c).value}")
    except Exception as e:
        print(f"Error {fpath}: {e}")
